#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cruzreta.py - Reconstrução 3D por Cruzamento de Retas (Biomecânica / Fotogrametria)
Disciplina: Biomecânica - Universidade de São Paulo (USP)

Implementa de forma intuitiva, modular e simplificada o método de cruzamento de retas
utilizado na planilha cruzreta.xlsx para estimar as coordenadas tridimensionais (X, Y, Z)
de um ponto no espaço a partir de duas câmeras (pontos focais) e seus respectivos pontos
de projeção nos planos ortogonais (XY, XZ, ZY).
"""

import sys
import os
import argparse
import json
import math
from dataclasses import dataclass, asdict
from typing import Optional, Tuple, Dict, Any, List

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import matplotlib
    # Use non-interactive backend by default if no display, but interactive if DISPLAY is set
    if not os.environ.get("DISPLAY"):
        matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from mpl_toolkits.mplot3d import Axes3D  # noqa: F401
except ImportError:
    plt = None


@dataclass
class Ponto3D:
    x: float
    y: float
    z: float

    def as_tuple(self) -> Tuple[float, float, float]:
        return (self.x, self.y, self.z)

    def __iter__(self):
        return iter((self.x, self.y, self.z))


@dataclass
class ResultadoPlano:
    nome: str
    eixo_h: str
    eixo_v: str
    mr1: float
    mr2: float
    br1: float
    br2: float
    coord_h: float
    coord_v: float


@dataclass
class ResultadoReconstrucao:
    c1: Ponto3D
    p1: Ponto3D
    c2: Ponto3D
    p2: Ponto3D
    plano_xy: ResultadoPlano
    plano_xz: ResultadoPlano
    plano_zy: ResultadoPlano
    rec3d: Ponto3D
    real: Optional[Ponto3D] = None
    erro_com_sinal: Optional[Ponto3D] = None
    erro_absoluto: Optional[Ponto3D] = None
    erro_total: Optional[float] = None
    erro_euclidiano: Optional[float] = None


class CruzamentoRetas:
    """
    Classe para cálculo de Reconstrução 3D por Cruzamento de Retas em Biomecânica.
    """

    DEFAULT_C1 = Ponto3D(365.8, 211.0, 107.0)
    DEFAULT_P1 = Ponto3D(0.0, 133.9, 50.9)
    DEFAULT_C2 = Ponto3D(131.4, 384.8, 151.8)
    DEFAULT_P2 = Ponto3D(179.4, 0.0, 17.1)
    DEFAULT_REAL = Ponto3D(159.3, 168.2, 74.9)

    def __init__(
        self,
        c1: Optional[Ponto3D] = None,
        p1: Optional[Ponto3D] = None,
        c2: Optional[Ponto3D] = None,
        p2: Optional[Ponto3D] = None,
        real: Optional[Ponto3D] = None,
    ):
        self.c1 = c1 or self.DEFAULT_C1
        self.p1 = p1 or self.DEFAULT_P1
        self.c2 = c2 or self.DEFAULT_C2
        self.p2 = p2 or self.DEFAULT_P2
        self.real = real if real is not None else self.DEFAULT_REAL

    @staticmethod
    def _calcular_intersecao_2d(
        nome: str,
        eixo_h: str,
        eixo_v: str,
        ch1: float,
        cv1: float,
        ph1: float,
        pv1: float,
        ch2: float,
        cv2: float,
        ph2: float,
        pv2: float,
    ) -> ResultadoPlano:
        """
        Calcula os coeficientes angulares (mr), lineares (br) e o ponto de interseção
        em um plano bidimensional entre as duas retas das câmeras.
        """
        delta_h1 = ph1 - ch1
        delta_h2 = ph2 - ch2

        if abs(delta_h1) < 1e-12:
            raise ValueError(f"No plano {nome}, a reta 1 é vertical (ph1 == ch1 == {ch1}).")
        if abs(delta_h2) < 1e-12:
            raise ValueError(f"No plano {nome}, a reta 2 é vertical (ph2 == ch2 == {ch2}).")

        mr1 = (pv1 - cv1) / delta_h1
        mr2 = (pv2 - cv2) / delta_h2

        # br1 = cv1 - mr1 * ch1 = -1 * (mr1 * ch1 - cv1)  (fórmula idêntica ao Excel)
        br1 = -1.0 * (mr1 * ch1 - cv1)
        br2 = -1.0 * (mr2 * ch2 - cv2)

        if abs(mr1 - mr2) < 1e-12:
            raise ValueError(f"No plano {nome}, as duas retas são paralelas (mr1 == mr2 == {mr1}).")

        # Ponto de interseção H (eixo horizontal) e V (eixo vertical)
        # H = (-(mr2 * ch2) + cv2 + (mr1 * ch1) - cv1) / (mr1 - mr2)
        coord_h = (-(mr2 * ch2) + cv2 + (mr1 * ch1) - cv1) / (mr1 - mr2)
        coord_v = (mr1 * coord_h) - (mr1 * ch1) + cv1

        return ResultadoPlano(
            nome=nome,
            eixo_h=eixo_h,
            eixo_v=eixo_v,
            mr1=mr1,
            mr2=mr2,
            br1=br1,
            br2=br2,
            coord_h=coord_h,
            coord_v=coord_v,
        )

    def resolver(self) -> ResultadoReconstrucao:
        """
        Executa a reconstrução 3D calculando o cruzamento nos planos XY, XZ e ZY.
        """
        # 1. Plano XY (horizontal: X, vertical: Y)
        plano_xy = self._calcular_intersecao_2d(
            nome="XY",
            eixo_h="X",
            eixo_v="Y",
            ch1=self.c1.x,
            cv1=self.c1.y,
            ph1=self.p1.x,
            pv1=self.p1.y,
            ch2=self.c2.x,
            cv2=self.c2.y,
            ph2=self.p2.x,
            pv2=self.p2.y,
        )

        # 2. Plano XZ (horizontal: X, vertical: Z)
        plano_xz = self._calcular_intersecao_2d(
            nome="XZ",
            eixo_h="X",
            eixo_v="Z",
            ch1=self.c1.x,
            cv1=self.c1.z,
            ph1=self.p1.x,
            pv1=self.p1.z,
            ch2=self.c2.x,
            cv2=self.c2.z,
            ph2=self.p2.x,
            pv2=self.p2.z,
        )

        # 3. Plano ZY (horizontal: Z, vertical: Y)
        plano_zy = self._calcular_intersecao_2d(
            nome="ZY",
            eixo_h="Z",
            eixo_v="Y",
            ch1=self.c1.z,
            cv1=self.c1.y,
            ph1=self.p1.z,
            pv1=self.p1.y,
            ch2=self.c2.z,
            cv2=self.c2.y,
            ph2=self.p2.z,
            pv2=self.p2.y,
        )

        # Reconstrução 3D pela média das estimativas dos pares de planos ortogonais:
        # X: média entre XY e XZ
        # Y: média entre XY e ZY
        # Z: média entre XZ e ZY
        rec_x = (plano_xy.coord_h + plano_xz.coord_h) / 2.0
        rec_y = (plano_xy.coord_v + plano_zy.coord_v) / 2.0
        rec_z = (plano_xz.coord_v + plano_zy.coord_h) / 2.0
        rec3d = Ponto3D(rec_x, rec_y, rec_z)

        # Cálculo de erros se coordenadas reais forem informadas
        erro_sinal = None
        erro_abs = None
        erro_total = None
        erro_eucl = None

        if self.real is not None:
            ex = self.real.x - rec_x
            ey = self.real.y - rec_y
            ez = self.real.z - rec_z
            erro_sinal = Ponto3D(ex, ey, ez)
            erro_abs = Ponto3D(abs(ex), abs(ey), abs(ez))
            erro_total = (abs(ex) + abs(ey) + abs(ez)) / 3.0
            erro_eucl = math.sqrt(ex**2 + ey**2 + ez**2)

        return ResultadoReconstrucao(
            c1=self.c1,
            p1=self.p1,
            c2=self.c2,
            p2=self.p2,
            plano_xy=plano_xy,
            plano_xz=plano_xz,
            plano_zy=plano_zy,
            rec3d=rec3d,
            real=self.real,
            erro_com_sinal=erro_sinal,
            erro_absoluto=erro_abs,
            erro_total=erro_total,
            erro_euclidiano=erro_eucl,
        )

    @classmethod
    def carregar_planilha_excel(cls, caminho_excel: str) -> "CruzamentoRetas":
        """
        Lê uma planilha compatível com cruzreta.xlsx e instancia a classe.
        """
        if not openpyxl:
            raise ImportError("Instale o openpyxl para carregar arquivos Excel: pip install openpyxl")

        wb = openpyxl.load_workbook(caminho_excel, data_only=True)
        ws = wb.active

        # Linha 2 = XY (xc1, yc1, xp1, yp1, xc2, yc2, xp2, yp2)
        # Linha 3 = XZ (xc1, zc1, xp1, zp1, xc2, zc2, xp2, zp2)
        # Linha 4 = ZY (zc1, yc1, zp1, yp1, zc2, yc2, zp2, yp2)
        # Linha 9 = Real (Q9=x, R9=y, S9=z)
        xc1 = float(ws.cell(2, 1).value)
        yc1 = float(ws.cell(2, 2).value)
        zc1 = float(ws.cell(3, 2).value)

        xp1 = float(ws.cell(2, 3).value)
        yp1 = float(ws.cell(2, 4).value)
        zp1 = float(ws.cell(3, 4).value)

        xc2 = float(ws.cell(2, 5).value)
        yc2 = float(ws.cell(2, 6).value)
        zc2 = float(ws.cell(3, 6).value)

        xp2 = float(ws.cell(2, 7).value)
        yp2 = float(ws.cell(2, 8).value)
        zp2 = float(ws.cell(3, 8).value)

        c1 = Ponto3D(xc1, yc1, zc1)
        p1 = Ponto3D(xp1, yp1, zp1)
        c2 = Ponto3D(xc2, yc2, zc2)
        p2 = Ponto3D(xp2, yp2, zp2)

        real = None
        q9 = ws.cell(9, 17).value
        r9 = ws.cell(9, 18).value
        s9 = ws.cell(9, 19).value
        if q9 is not None and r9 is not None and s9 is not None:
            real = Ponto3D(float(q9), float(r9), float(s9))

        return cls(c1=c1, p1=p1, c2=c2, p2=p2, real=real)


def imprimir_relatorio(resultado: ResultadoReconstrucao):
    """
    Exibe um relatório visual, limpo e detalhado no terminal.
    """
    sep = "=" * 76
    sub_sep = "-" * 76

    print(sep)
    print("      RECONSTRUÇÃO 3D POR CRUZAMENTO DE RETAS - BIOMECÂNICA (USP)     ")
    print(sep)

    print("\n[1] PARÂMETROS DE ENTRADA (COORDENADAS 3D):")
    print(f"  • Câmera 1 (Focal Center C1):  X = {resultado.c1.x:8.2f} | Y = {resultado.c1.y:8.2f} | Z = {resultado.c1.z:8.2f}")
    print(f"  • Ponto Projeção 1 (P1):       X = {resultado.p1.x:8.2f} | Y = {resultado.p1.y:8.2f} | Z = {resultado.p1.z:8.2f}")
    print(f"  • Câmera 2 (Focal Center C2):  X = {resultado.c2.x:8.2f} | Y = {resultado.c2.y:8.2f} | Z = {resultado.c2.z:8.2f}")
    print(f"  • Ponto Projeção 2 (P2):       X = {resultado.p2.x:8.2f} | Y = {resultado.p2.y:8.2f} | Z = {resultado.p2.z:8.2f}")

    print("\n[2] EQUAÇÕES DE RETAS NOS PLANOS ORTOGONAIS (Y = mr * X + br):")
    print(sub_sep)
    print(f"{'Plano':<8} | {'mr1 (slope 1)':<14} | {'br1 (interc 1)':<14} | {'mr2 (slope 2)':<14} | {'br2 (interc 2)':<14}")
    print(sub_sep)

    for p in [resultado.plano_xy, resultado.plano_xz, resultado.plano_zy]:
        print(f"{p.nome:<8} | {p.mr1:14.6f} | {p.br1:14.4f} | {p.mr2:14.6f} | {p.br2:14.4f}")
    print(sub_sep)

    print("\n[3] PONTOS DE INTERSEÇÃO 2D EM CADA PLANO:")
    print(sub_sep)
    print(f"  • Plano XY ->  X = {resultado.plano_xy.coord_h:10.4f}  |  Y = {resultado.plano_xy.coord_v:10.4f}")
    print(f"  • Plano XZ ->  X = {resultado.plano_xz.coord_h:10.4f}  |  Z = {resultado.plano_xz.coord_v:10.4f}")
    print(f"  • Plano ZY ->  Z = {resultado.plano_zy.coord_h:10.4f}  |  Y = {resultado.plano_zy.coord_v:10.4f}")
    print(sub_sep)

    print("\n[4] RECONSTRUÇÃO 3D FINAL (REC3D):")
    print(f"  • X = Média(X_xy, X_xz) = ({resultado.plano_xy.coord_h:.4f} + {resultado.plano_xz.coord_h:.4f}) / 2 = {resultado.rec3d.x:10.4f}")
    print(f"  • Y = Média(Y_xy, Y_zy) = ({resultado.plano_xy.coord_v:.4f} + {resultado.plano_zy.coord_v:.4f}) / 2 = {resultado.rec3d.y:10.4f}")
    print(f"  • Z = Média(Z_xz, Z_zy) = ({resultado.plano_xz.coord_v:.4f} + {resultado.plano_zy.coord_h:.4f}) / 2 = {resultado.rec3d.z:10.4f}")

    if resultado.real and resultado.erro_com_sinal and resultado.erro_absoluto:
        print("\n[5] ANÁLISE DE EXATIDÃO (COMPARAÇÃO COM COORDENADAS REAIS):")
        print(sub_sep)
        print(f"{'Eixo':<6} | {'Reconstruído':<14} | {'Real (Gabarito)':<16} | {'Erro (Real - Rec)':<18} | {'Erro Absoluto':<14}")
        print(sub_sep)
        print(f"{'X':<6} | {resultado.rec3d.x:14.4f} | {resultado.real.x:16.4f} | {resultado.erro_com_sinal.x:18.4f} | {resultado.erro_absoluto.x:14.4f}")
        print(f"{'Y':<6} | {resultado.rec3d.y:14.4f} | {resultado.real.y:16.4f} | {resultado.erro_com_sinal.y:18.4f} | {resultado.erro_absoluto.y:14.4f}")
        print(f"{'Z':<6} | {resultado.rec3d.z:14.4f} | {resultado.real.z:16.4f} | {resultado.erro_com_sinal.z:18.4f} | {resultado.erro_absoluto.z:14.4f}")
        print(sub_sep)
        print(f"  • Erro Médio Absoluto (Total Error da Planilha): {resultado.erro_total:8.4f}")
        print(f"  • Erro Euclidiano 3D (Distância Real-Rec):       {resultado.erro_euclidiano:8.4f}")

    print(sep + "\n")


def plotar_reconstrucao(resultado: ResultadoReconstrucao, caminho_salvar: Optional[str] = None, mostrar: bool = True):
    """
    Gera um gráfico visual com 4 painéis:
    - Painel 1: Visualização 3D do espaço, câmeras, retas de visão e ponto reconstruído.
    - Painel 2: Cruzamento no Plano XY.
    - Painel 3: Cruzamento no Plano XZ.
    - Painel 4: Cruzamento no Plano ZY.
    """
    if plt is None:
        print("Aviso: matplotlib não está instalado. Não foi possível gerar os gráficos.")
        return

    fig = plt.figure(figsize=(15, 11))
    fig.suptitle("Biomecânica USP - Reconstrução 3D por Cruzamento de Retas", fontsize=15, fontweight="bold")

    # 1. Painel 3D
    ax3d = fig.add_subplot(2, 2, 1, projection="3d")
    ax3d.set_title("Espaço Tridimensional (3D)", fontweight="bold")

    # Posições das câmeras e projeções
    ax3d.scatter(resultado.c1.x, resultado.c1.y, resultado.c1.z, color="blue", s=100, label="Câmera 1 (C1)", marker="^")
    ax3d.scatter(resultado.p1.x, resultado.p1.y, resultado.p1.z, color="cyan", s=70, label="Projeção 1 (P1)", marker="o")
    ax3d.scatter(resultado.c2.x, resultado.c2.y, resultado.c2.z, color="red", s=100, label="Câmera 2 (C2)", marker="^")
    ax3d.scatter(resultado.p2.x, resultado.p2.y, resultado.p2.z, color="orange", s=70, label="Projeção 2 (P2)", marker="o")

    # Retas 3D das câmeras passando pelo ponto reconstruído
    ax3d.plot([resultado.c1.x, resultado.p1.x], [resultado.c1.y, resultado.p1.y], [resultado.c1.z, resultado.p1.z], "b--", alpha=0.6, label="Raio Óptico 1")
    ax3d.plot([resultado.c2.x, resultado.p2.x], [resultado.c2.y, resultado.p2.y], [resultado.c2.z, resultado.p2.z], "r--", alpha=0.6, label="Raio Óptico 2")

    # Ponto Reconstruído
    ax3d.scatter(resultado.rec3d.x, resultado.rec3d.y, resultado.rec3d.z, color="green", s=140, label="REC3D Reconstruído", marker="*")

    # Ponto Real (se houver)
    if resultado.real:
        ax3d.scatter(resultado.real.x, resultado.real.y, resultado.real.z, color="purple", s=110, label="Ponto Real", marker="x")
        # Linha de erro entre reconstruído e real
        ax3d.plot(
            [resultado.rec3d.x, resultado.real.x],
            [resultado.rec3d.y, resultado.real.y],
            [resultado.rec3d.z, resultado.real.z],
            "m-.",
            linewidth=1.5,
            label=f"Erro Euclidiano ({resultado.erro_euclidiano:.2f})"
        )

    ax3d.set_xlabel("X")
    ax3d.set_ylabel("Y")
    ax3d.set_zlabel("Z")
    ax3d.legend(loc="upper left", fontsize=8)
    ax3d.grid(True, alpha=0.3)

    # Função auxiliar para plot 2D
    def plot_2d_plane(ax, plano: ResultadoPlano, c1_h, c1_v, p1_h, p1_v, c2_h, c2_v, p2_h, p2_v, real_h=None, real_v=None):
        ax.set_title(f"Plano {plano.nome} ({plano.eixo_h} x {plano.eixo_v})", fontweight="bold")
        ax.scatter(c1_h, c1_v, color="blue", marker="^", s=80, label="C1")
        ax.scatter(p1_h, p1_v, color="cyan", marker="o", s=60, label="P1")
        ax.scatter(c2_h, c2_v, color="red", marker="^", s=80, label="C2")
        ax.scatter(p2_h, p2_v, color="orange", marker="o", s=60, label="P2")

        # Gerar faixa de pontos para desenhar as retas
        all_h = [c1_h, p1_h, c2_h, p2_h, plano.coord_h]
        if real_h is not None:
            all_h.append(real_h)
        min_h, max_h = min(all_h) - 20, max(all_h) + 20
        h_vals = [min_h, max_h]

        v_line1 = [plano.mr1 * h + plano.br1 for h in h_vals]
        v_line2 = [plano.mr2 * h + plano.br2 for h in h_vals]

        ax.plot(h_vals, v_line1, "b--", alpha=0.6, label="Reta Câmera 1")
        ax.plot(h_vals, v_line2, "r--", alpha=0.6, label="Reta Câmera 2")

        # Ponto de cruzamento
        ax.scatter(plano.coord_h, plano.coord_v, color="green", marker="*", s=130, zorder=5,
                   label=f"Interseção ({plano.coord_h:.1f}, {plano.coord_v:.1f})")

        if real_h is not None and real_v is not None:
            ax.scatter(real_h, real_v, color="purple", marker="x", s=90, zorder=5,
                       label=f"Real ({real_h:.1f}, {real_v:.1f})")

        ax.set_xlabel(plano.eixo_h)
        ax.set_ylabel(plano.eixo_v)
        ax.grid(True, linestyle=":", alpha=0.5)
        ax.legend(fontsize=8)

    # 2. Painel Plano XY
    ax_xy = fig.add_subplot(2, 2, 2)
    real_x = resultado.real.x if resultado.real else None
    real_y = resultado.real.y if resultado.real else None
    plot_2d_plane(
        ax_xy, resultado.plano_xy,
        resultado.c1.x, resultado.c1.y, resultado.p1.x, resultado.p1.y,
        resultado.c2.x, resultado.c2.y, resultado.p2.x, resultado.p2.y,
        real_x, real_y
    )

    # 3. Painel Plano XZ
    ax_xz = fig.add_subplot(2, 2, 3)
    real_z = resultado.real.z if resultado.real else None
    plot_2d_plane(
        ax_xz, resultado.plano_xz,
        resultado.c1.x, resultado.c1.z, resultado.p1.x, resultado.p1.z,
        resultado.c2.x, resultado.c2.z, resultado.p2.x, resultado.p2.z,
        real_x, real_z
    )

    # 4. Painel Plano ZY
    ax_zy = fig.add_subplot(2, 2, 4)
    plot_2d_plane(
        ax_zy, resultado.plano_zy,
        resultado.c1.z, resultado.c1.y, resultado.p1.z, resultado.p1.y,
        resultado.c2.z, resultado.c2.y, resultado.p2.z, resultado.p2.y,
        real_z, real_y
    )

    plt.tight_layout()

    if caminho_salvar:
        plt.savefig(caminho_salvar, dpi=200, bbox_inches="tight")
        print(f" Gráfico salvo com sucesso em: {caminho_salvar}")

    if mostrar and os.environ.get("DISPLAY"):
        plt.show()
    plt.close()


def exportar_para_excel(resultado: ResultadoReconstrucao, caminho_saida: str):
    """
    Exporta os cálculos para um arquivo Excel com estrutura compatível e limpa.
    """
    if not openpyxl:
        print("Erro: openpyxl necessário para exportar em formato .xlsx.")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reconstrucao3D"

    # Cabeçalhos originais
    headers = ["Plano", "ch1", "cv1", "ph1", "pv1", "ch2", "cv2", "ph2", "pv2", "mr1", "mr2", "br1", "br2", "Coord_H", "Coord_V"]
    ws.append(headers)

    planos_data = [
        ("XY", resultado.c1.x, resultado.c1.y, resultado.p1.x, resultado.p1.y,
         resultado.c2.x, resultado.c2.y, resultado.p2.x, resultado.p2.y,
         resultado.plano_xy.mr1, resultado.plano_xy.mr2, resultado.plano_xy.br1, resultado.plano_xy.br2,
         resultado.plano_xy.coord_h, resultado.plano_xy.coord_v),
        ("XZ", resultado.c1.x, resultado.c1.z, resultado.p1.x, resultado.p1.z,
         resultado.c2.x, resultado.c2.z, resultado.p2.x, resultado.p2.z,
         resultado.plano_xz.mr1, resultado.plano_xz.mr2, resultado.plano_xz.br1, resultado.plano_xz.br2,
         resultado.plano_xz.coord_h, resultado.plano_xz.coord_v),
        ("ZY", resultado.c1.z, resultado.c1.y, resultado.p1.z, resultado.p1.y,
         resultado.c2.z, resultado.c2.y, resultado.p2.z, resultado.p2.y,
         resultado.plano_zy.mr1, resultado.plano_zy.mr2, resultado.plano_zy.br1, resultado.plano_zy.br2,
         resultado.plano_zy.coord_h, resultado.plano_zy.coord_v),
    ]

    for p in planos_data:
        ws.append(list(p))

    ws.append([])
    ws.append(["Métrica", "X", "Y", "Z"])
    ws.append(["REC3D (Reconstruído)", resultado.rec3d.x, resultado.rec3d.y, resultado.rec3d.z])
    if resultado.real:
        ws.append(["Real (Gabarito)", resultado.real.x, resultado.real.y, resultado.real.z])
        ws.append(["Erro (Real - Rec)", resultado.erro_com_sinal.x, resultado.erro_com_sinal.y, resultado.erro_com_sinal.z])
        ws.append(["Erro Absoluto", resultado.erro_absoluto.x, resultado.erro_absoluto.y, resultado.erro_absoluto.z])
        ws.append(["Erro Total (Média)", resultado.erro_total, "", ""])
        ws.append(["Erro Euclidiano", resultado.erro_euclidiano, "", ""])

    wb.save(caminho_saida)
    print(f" Resultados exportados para o Excel em: {caminho_saida}")


def solicitar_ponto_interativo(nome: str, default: Ponto3D) -> Ponto3D:
    """
    Solicita as coordenadas 3D de um ponto no terminal com valor default intuitivo.
    """
    prompt = f"  • {nome} [padrão: {default.x} {default.y} {default.z}]: "
    try:
        entrada = input(prompt).strip()
    except (EOFError, KeyboardInterrupt):
        print("\nOperação cancelada pelo usuário.")
        sys.exit(0)

    if not entrada:
        return default

    partes = entrada.replace(",", ".").split()
    if len(partes) != 3:
        print(f"    [!] Formato inválido. Usando valor padrão: {default.x} {default.y} {default.z}")
        return default

    try:
        x, y, z = float(partes[0]), float(partes[1]), float(partes[2])
        return Ponto3D(x, y, z)
    except ValueError:
        print(f"    [!] Valores numéricos inválidos. Usando padrão.")
        return default


def modo_interativo() -> CruzamentoRetas:
    """
    Executa assistente de entrada interativa no terminal.
    """
    print("\n" + "=" * 76)
    print("      ASSISTENTE INTERATIVO - RECONSTRUÇÃO 3D POR CRUZAMENTO DE RETAS")
    print("=" * 76)
    print("Pressione [ENTER] para aceitar os valores padrão da tarefa (cruzreta.xlsx),\n"
          "ou digite as 3 coordenadas separadas por espaço (ex: 365.8 211.0 107.0).\n")

    c1 = solicitar_ponto_interativo("Câmera 1 (Centro Focal C1)", CruzamentoRetas.DEFAULT_C1)
    p1 = solicitar_ponto_interativo("Ponto de Projeção 1 (P1)", CruzamentoRetas.DEFAULT_P1)
    c2 = solicitar_ponto_interativo("Câmera 2 (Centro Focal C2)", CruzamentoRetas.DEFAULT_C2)
    p2 = solicitar_ponto_interativo("Ponto de Projeção 2 (P2)", CruzamentoRetas.DEFAULT_P2)
    real = solicitar_ponto_interativo("Ponto Real / Gabarito", CruzamentoRetas.DEFAULT_REAL)

    return CruzamentoRetas(c1=c1, p1=p1, c2=c2, p2=p2, real=real)


def main():
    parser = argparse.ArgumentParser(
        description="Reconstrução 3D por Cruzamento de Retas (Biomecânica USP - cruzreta.xlsx)"
    )
    parser.add_argument(
        "-i", "--interativo",
        action="store_true",
        help="Executa em modo interativo guiado no terminal"
    )
    parser.add_argument(
        "-a", "--arquivo",
        type=str,
        default=None,
        help="Caminho para carregar planilha Excel (.xlsx) existente"
    )
    parser.add_argument(
        "-p", "--plot",
        action="store_true",
        help="Exibe janela com visualização gráfica dos cruzamentos e 3D"
    )
    parser.add_argument(
        "-g", "--salvar-grafico",
        type=str,
        default=None,
        nargs="?",
        const="cruzreta_reconstrucao3d.png",
        help="Caminho para salvar imagem gráfica PNG (padrão: cruzreta_reconstrucao3d.png)"
    )
    parser.add_argument(
        "-w", "--web",
        action="store_true",
        help="Inicia interface web visual e interativa no navegador"
    )
    parser.add_argument(
        "-e", "--exportar",
        type=str,
        default=None,
        help="Exporta resultados calculados para planilha Excel (.xlsx)"
    )
    parser.add_argument(
        "-j", "--json",
        action="store_true",
        help="Emite a saída dos cálculos em formato JSON estruturado"
    )

    args = parser.parse_args()

    if args.web:
        try:
            from web_app import run_web_server
            run_web_server()
            return
        except ImportError:
            print("Erro ao carregar módulo web_app. Certifique-se de que web_app.py está no mesmo diretório.")
            sys.exit(1)

    # Seleção da fonte dos dados
    if args.interativo:
        calculador = modo_interativo()
    elif args.arquivo:
        if not os.path.exists(args.arquivo):
            print(f"Erro: Arquivo '{args.arquivo}' não encontrado.")
            sys.exit(1)
        calculador = CruzamentoRetas.carregar_planilha_excel(args.arquivo)
    else:
        # Se existir cruzreta.xlsx no diretório atual ou padrão, tenta carregar dele para fidelidade total
        caminhos_busca = [
            "cruzreta.xlsx",
            "/home/luis/Documents/USP/biomecanica/tarefa 1/cruzreta.xlsx",
        ]
        carregado = False
        for c in caminhos_busca:
            if os.path.exists(c):
                try:
                    calculador = CruzamentoRetas.carregar_planilha_excel(c)
                    carregado = True
                    break
                except Exception:
                    pass
        if not carregado:
            calculador = CruzamentoRetas()

    resultado = calculador.resolver()

    if args.json:
        # Serializar dataclass para JSON
        def serializer(obj):
            if hasattr(obj, "__dict__"):
                return asdict(obj)
            return str(obj)
        print(json.dumps(asdict(resultado), indent=2, default=serializer))
    else:
        imprimir_relatorio(resultado)

    if args.exportar:
        exportar_para_excel(resultado, args.exportar)

    if args.salvar_grafico or args.plot:
        caminho_img = args.salvar_grafico or "cruzreta_reconstrucao3d.png"
        if not os.path.isabs(caminho_img) and not os.path.dirname(caminho_img):
            script_dir = os.path.dirname(os.path.abspath(__file__))
            caminho_img = os.path.join(script_dir, caminho_img)
        plotar_reconstrucao(resultado, caminho_salvar=caminho_img, mostrar=args.plot)


if __name__ == "__main__":
    main()
