#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_cruzreta.py - Testes unitários para o script de reconstrução 3D por cruzamento de retas.
Compara valores calculados contra a planilha de referência cruzreta.xlsx.
"""

import os
import sys
import unittest
import openpyxl

from cruzreta import CruzamentoRetas, Ponto3D, ResultadoReconstrucao


class TestCruzamentoRetas(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cruzreta.xlsx")
        cls.wb = openpyxl.load_workbook(cls.excel_path, data_only=True)
        cls.ws = cls.wb["Sheet1"]

    def test_valores_exatos_planilha_cruzreta(self):
        """
        Verifica se todos os cálculos (coeficientes, interseções, REC3D e erros)
        coincidem com os valores da planilha cruzreta.xlsx até 6 casas decimais.
        """
        calc = CruzamentoRetas.carregar_planilha_excel(self.excel_path)
        res = calc.resolver()

        ws = self.ws

        # 1. Plano XY (Linha 2 do Excel)
        self.assertAlmostEqual(res.plano_xy.mr1, float(ws["J2"].value), places=6)
        self.assertAlmostEqual(res.plano_xy.mr2, float(ws["K2"].value), places=6)
        self.assertAlmostEqual(res.plano_xy.br1, float(ws["L2"].value), places=6)
        self.assertAlmostEqual(res.plano_xy.br2, float(ws["M2"].value), places=6)
        self.assertAlmostEqual(res.plano_xy.coord_h, float(ws["O2"].value), places=6)
        self.assertAlmostEqual(res.plano_xy.coord_v, float(ws["P2"].value), places=6)

        # 2. Plano XZ (Linha 3 do Excel)
        self.assertAlmostEqual(res.plano_xz.mr1, float(ws["J3"].value), places=6)
        self.assertAlmostEqual(res.plano_xz.mr2, float(ws["K3"].value), places=6)
        self.assertAlmostEqual(res.plano_xz.br1, float(ws["L3"].value), places=6)
        self.assertAlmostEqual(res.plano_xz.br2, float(ws["M3"].value), places=6)
        self.assertAlmostEqual(res.plano_xz.coord_h, float(ws["O3"].value), places=6)
        self.assertAlmostEqual(res.plano_xz.coord_v, float(ws["P3"].value), places=6)

        # 3. Plano ZY (Linha 4 do Excel)
        self.assertAlmostEqual(res.plano_zy.mr1, float(ws["J4"].value), places=6)
        self.assertAlmostEqual(res.plano_zy.mr2, float(ws["K4"].value), places=6)
        self.assertAlmostEqual(res.plano_zy.br1, float(ws["L4"].value), places=6)
        self.assertAlmostEqual(res.plano_zy.br2, float(ws["M4"].value), places=6)
        self.assertAlmostEqual(res.plano_zy.coord_h, float(ws["O4"].value), places=6)
        self.assertAlmostEqual(res.plano_zy.coord_v, float(ws["P4"].value), places=6)

        # 4. REC3D (Linha 8 do Excel: Q8=x, R8=y, S8=z)
        self.assertAlmostEqual(res.rec3d.x, float(ws["Q8"].value), places=6)
        self.assertAlmostEqual(res.rec3d.y, float(ws["R8"].value), places=6)
        self.assertAlmostEqual(res.rec3d.z, float(ws["S8"].value), places=6)

        # 5. Coordenadas Reais (Linha 9: Q9=159.3, R9=168.2, S9=74.9)
        self.assertAlmostEqual(res.real.x, float(ws["Q9"].value), places=6)
        self.assertAlmostEqual(res.real.y, float(ws["R9"].value), places=6)
        self.assertAlmostEqual(res.real.z, float(ws["S9"].value), places=6)

        # 6. Erro com sinal (Linha 10)
        self.assertAlmostEqual(res.erro_com_sinal.x, float(ws["Q10"].value), places=6)
        self.assertAlmostEqual(res.erro_com_sinal.y, float(ws["R10"].value), places=6)
        self.assertAlmostEqual(res.erro_com_sinal.z, float(ws["S10"].value), places=6)

        # 7. Erro absoluto (Linha 11)
        self.assertAlmostEqual(res.erro_absoluto.x, float(ws["Q11"].value), places=6)
        self.assertAlmostEqual(res.erro_absoluto.y, float(ws["R11"].value), places=6)
        self.assertAlmostEqual(res.erro_absoluto.z, float(ws["S11"].value), places=6)

        # 8. Erro total médio (Linha 12: Q12)
        self.assertAlmostEqual(res.erro_total, float(ws["Q12"].value), places=6)

    def test_linhas_paralelas_lanca_excecao(self):
        """Verifica se retas paralelas disparam erro tratado."""
        c1 = Ponto3D(0.0, 0.0, 0.0)
        p1 = Ponto3D(1.0, 1.0, 1.0)
        c2 = Ponto3D(0.0, 1.0, 0.0)
        p2 = Ponto3D(1.0, 2.0, 1.0)  # Mesma inclinação nos eixos

        calc = CruzamentoRetas(c1, p1, c2, p2)
        with self.assertRaises(ValueError):
            calc.resolver()


if __name__ == "__main__":
    unittest.main()
